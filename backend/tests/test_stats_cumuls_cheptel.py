"""Calculs cumulés (émis / contrôlés), cheptel par pays-année-mois, filtre par mois, export Excel."""
import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.models.controle import Controle, ModeVerification, ResultatControle
from app.models.passeport import StatutPasseport
from tests.test_indicateurs_cheptel import _troupeau
from tests.test_statistiques import _commande_payee_avec_paiement


async def _controle(db, passeport, agent, resultat):
    db.add(Controle(passeport_id=passeport.id, poste_id="P1", agent_id=agent.id, resultat=resultat,
                    itineraire_disponible_localement=False, mode=ModeVerification.EN_LIGNE))
    passeport.statut = StatutPasseport.CONTROLE  # comportement réel de POST /controles
    await db.commit()


@pytest.mark.asyncio
async def test_emis_cumule_et_controles_distincts(client, db, super_admin, agent_controle_cmr, pays_cameroun, pays_tchad):
    u, h = super_admin
    ag, _ = agent_controle_cmr
    _, pp = await _commande_payee_avec_paiement(db, pays_cameroun.id, u.id, quantite=3)
    for p in pp:
        await _troupeau(db, p, [("bovin", 5)], pays_tchad.id, pays_cameroun.id)
    # p0 : scanné 3 fois (valide, refuse, a_verifier) ; p1 : 1 fois ; p2 jamais.
    await _controle(db, pp[0], ag, ResultatControle.VALIDE)
    await _controle(db, pp[0], ag, ResultatControle.REFUSE)
    await _controle(db, pp[0], ag, ResultatControle.A_VERIFIER)
    await _controle(db, pp[1], ag, ResultatControle.VALIDE)
    pp[2].statut = StatutPasseport.EMIS
    await db.commit()

    lignes = (await client.get("/api/v1/statistiques/par-pays-annee", headers=h)).json()
    cmr = next(l for l in lignes if l["pays_id"] == pays_cameroun.id)
    # 3 passeports émis (cumul), même si deux ont depuis le statut « contrôlé » ; 2 passeports contrôlés, chacun compté UNE fois (premier scan).
    assert cmr["passeports_cumules"]["emis"] == 3
    assert cmr["passeports_cumules"]["controle"] == 2
    # Un passeport scanné plusieurs fois ne compte qu'une fois : la répartition par résultat porte sur les 2 premiers scans.
    assert cmr["nb_controles"] == 2
    # Cohérence demandée : contrôlés = validés + à vérifier + refusés.
    assert cmr["passeports_cumules"]["controle"] == sum(cmr["controles_par_resultat"].values())

    bord = (await client.get("/api/v1/statistiques/tableau-bord", headers=h)).json()
    phases = {p["statut"]: p["nombre"] for p in bord["entonnoir_global"]}
    assert phases["emis"] == 3 and phases["controle"] == 2 and "precharge" not in phases


@pytest.mark.asyncio
async def test_cheptel_par_pays_annee_mois_et_filtre(client, db, super_admin, pays_cameroun, pays_tchad):
    u, h = super_admin
    _, pp = await _commande_payee_avec_paiement(db, pays_cameroun.id, u.id, quantite=2)
    await _troupeau(db, pp[0], [("bovin", 30), ("ovin", 10)], pays_tchad.id, pays_cameroun.id, vaccins=["PPR"])
    await _troupeau(db, pp[1], [("bovin", 20)], pays_tchad.id, pays_cameroun.id)
    maintenant = datetime.now(timezone.utc)
    r = (await client.get("/api/v1/statistiques/cheptel-pays-annee", headers=h)).json()
    assert len(r) == 1
    ligne = r[0]
    assert ligne["pays_id"] == pays_cameroun.id and ligne["annee"] == maintenant.year and ligne["mois"] == maintenant.month
    assert ligne["nb_passeports"] == 2 and ligne["total"] == 60
    assert ligne["par_espece"] == {"bovin": 50, "ovin": 10} and ligne["vaccines"] == 40
    autre_mois = 1 if maintenant.month != 1 else 2
    assert (await client.get(f"/api/v1/statistiques/cheptel-pays-annee?mois={autre_mois}", headers=h)).json() == []
    assert len((await client.get(f"/api/v1/statistiques/cheptel-pays-annee?mois={maintenant.month}&annee={maintenant.year}", headers=h)).json()) == 1
    assert (await client.get("/api/v1/statistiques/cheptel-pays-annee?base=x", headers=h)).status_code == 422
    # Filtre par mois aussi sur les commandes.
    cmd = (await client.get(f"/api/v1/statistiques/par-pays-annee?mois={maintenant.month}", headers=h)).json()
    assert any(l["nb_commandes"] > 0 for l in cmd)
    vide = (await client.get(f"/api/v1/statistiques/par-pays-annee?mois={autre_mois}", headers=h)).json()
    assert all(l["nb_commandes"] == 0 for l in vide)


@pytest.mark.asyncio
async def test_export_excel_cheptel_et_mois(client, db, super_admin, pays_cameroun, pays_tchad):
    u, h = super_admin
    _, pp = await _commande_payee_avec_paiement(db, pays_cameroun.id, u.id, quantite=1)
    await _troupeau(db, pp[0], [("bovin", 12)], pays_tchad.id, pays_cameroun.id)
    maintenant = datetime.now(timezone.utc)
    r = await client.get(f"/api/v1/statistiques/export?categories=cheptel,commandes&mois={maintenant.month}", headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Cheptel vérifié (émission)" in wb.sheetnames
    feuille = wb["Cheptel vérifié (émission)"]
    assert [c.value for c in feuille[2]][:5][3:] == [1, 12]
    assert (await client.get("/api/v1/statistiques/export?categories=cheptel&mois=13", headers=h)).status_code == 422


@pytest.mark.asyncio
async def test_indicateurs_filtres_par_mois(client, db, super_admin, pays_cameroun, pays_tchad):
    u, h = super_admin
    _, pp = await _commande_payee_avec_paiement(db, pays_cameroun.id, u.id, quantite=1)
    await _troupeau(db, pp[0], [("bovin", 7)], pays_tchad.id, pays_cameroun.id)
    m = datetime.now(timezone.utc)
    assert (await client.get(f"/api/v1/statistiques/indicateurs?mois={m.month}&annee={m.year}", headers=h)).json()["total_cheptel"] == 7
    autre = 1 if m.month != 1 else 2
    assert (await client.get(f"/api/v1/statistiques/indicateurs?mois={autre}", headers=h)).json()["total_cheptel"] == 0


@pytest.mark.asyncio
async def test_filtre_intervalle_dates_sur_les_statistiques(client, db, super_admin, pays_cameroun, pays_tchad):
    u, h = super_admin
    _, pp = await _commande_payee_avec_paiement(db, pays_cameroun.id, u.id, quantite=1)
    await _troupeau(db, pp[0], [("bovin", 12)], pays_tchad.id, pays_cameroun.id)
    aujourdhui = datetime.now(timezone.utc).date()
    hier = aujourdhui.replace(day=1) if aujourdhui.day > 1 else aujourdhui
    # Intervalle contenant aujourd'hui : données présentes.
    r = (await client.get("/api/v1/statistiques/indicateurs", params={"debut": str(hier), "fin": str(aujourdhui)}, headers=h)).json()
    assert r["total_cheptel"] == 12
    # Intervalle dans le passé : plus rien.
    r = (await client.get("/api/v1/statistiques/indicateurs", params={"debut": "2020-01-01", "fin": "2020-12-31"}, headers=h)).json()
    assert r["total_cheptel"] == 0
    r = (await client.get("/api/v1/statistiques/par-pays-annee", params={"debut": "2020-01-01", "fin": "2020-12-31"}, headers=h)).json()
    assert all(l["nb_commandes"] == 0 for l in r)
    r = (await client.get("/api/v1/statistiques/par-pays", params={"debut": "2020-01-01", "fin": "2020-12-31"}, headers=h)).json()
    assert all(l["nb_commandes"] == 0 and l["montant_encaisse_xaf"] == 0 for l in r)
    r = (await client.get("/api/v1/statistiques/cheptel-pays-annee", params={"debut": "2020-01-01", "fin": "2020-12-31"}, headers=h)).json()
    assert r == []
    ok = await client.get("/api/v1/statistiques/export", params={"categories": "commandes", "debut": "2020-01-01", "fin": "2020-12-31"}, headers=h)
    assert ok.status_code == 200
